"""2026-09-27 — company documents in fleet chat (app/company_docs.py,
app/routers/company_documents.py)."""
import asyncio
import importlib
import io
import uuid
from datetime import datetime, timezone

import pytest
from fastapi import HTTPException, UploadFile

from app import company_docs as CD
from app.auth.schemas import CurrentUser

R = importlib.import_module("app.routers.company_documents")

USER = CurrentUser(user_id=str(uuid.uuid4()), email="dpa@fleet.example", role="captain", tier="captain")
WS = uuid.uuid4()

MANUAL = """COMPANY SAFETY MANAGEMENT SYSTEM

1 Purpose
This manual sets out how the fleet operates safely.

4.2 Emergency Drills
Fire and abandon-vessel drills are held monthly on every boat.
Each drill is logged in the TVR with the names of all participants.

4.3 Man Overboard
Recover the person using the rescue sling. Report to the office within one hour.
"""


def test_split_sections_at_headings():
    secs = CD.split_sections([MANUAL])
    heads = [s.heading for s in secs]
    # the ALL-CAPS title has no body of its own, so it doesn't become a section
    assert heads == ["1 Purpose", "4.2 Emergency Drills", "4.3 Man Overboard"]
    drills = next(s for s in secs if s.heading == "4.2 Emergency Drills")
    assert "held monthly" in drills.text and "rescue sling" not in drills.text


def test_without_headings_each_page_is_a_section():
    secs = CD.split_sections(["plain text on page one.", "", "more text on page three."])
    assert [s.heading for s in secs] == ["p.1", "p.3"]


def test_chunks_stay_within_the_token_budget():
    long = "\n\n".join(f"Paragraph {i}: " + "crew shall wear lifejackets on deck. " * 20 for i in range(12))
    chunks = CD.chunk_sections([CD.Section("5 PPE", long)], max_tokens=200)
    assert len(chunks) > 1 and all(h == "5 PPE" for h, _, _ in chunks)
    assert [i for _, i, _ in chunks] == list(range(len(chunks)))
    assert all(len(CD._ENC.encode(t)) <= 200 * 2 for _, _, t in chunks)


def test_label_round_trip_and_title():
    lab = CD.label("TSMS Manual", "4.2 Emergency Drills")
    assert lab == "Company: TSMS Manual §4.2 Emergency Drills"
    assert CD.parse_label(f"[{lab}]") == ("TSMS Manual", "4.2 Emergency Drills")
    assert CD.parse_label("46 CFR 140.420") is None
    assert CD.title_from_filename("Fleet_TSMS_Manual_rev3.pdf") == "Fleet TSMS Manual rev3"


def test_extract_text_docx_and_scanned_pdf(tmp_path):
    import docx
    from pypdf import PdfWriter
    d = docx.Document()
    d.add_paragraph("4.2 Emergency Drills")
    d.add_paragraph("Drills are monthly.")
    p = tmp_path / "m.docx"
    d.save(p)
    pages = CD.extract_pages(p, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert "Drills are monthly." in pages[0]
    w = PdfWriter()
    w.add_blank_page(width=612, height=792)
    pdf = tmp_path / "scan.pdf"
    with open(pdf, "wb") as f:
        w.write(f)
    with pytest.raises(CD.UnreadableDocument, match="no text layer"):
        CD.extract_pages(pdf, "application/pdf")


class _Pool:
    def __init__(self, fetchval=(), fetchrow=None, fetch=()):
        self.calls, self._fv, self._fr, self._f = [], list(fetchval), fetchrow, list(fetch)

    async def fetchval(self, sql, *args):
        self.calls.append(("fetchval", sql, args))
        return self._fv.pop(0) if self._fv else None

    async def fetchrow(self, sql, *args):
        self.calls.append(("fetchrow", sql, args))
        return self._fr

    async def fetch(self, sql, *args):
        self.calls.append(("fetch", sql, args))
        return self._f


def test_search_and_lookup_are_scoped_to_the_workspace():
    pool = _Pool(fetch=[{"id": 1, "section": "4.2 Emergency Drills", "chunk_index": 0, "text": "monthly", "title": "TSMS"}])
    rows = asyncio.run(CD.search(pool, WS, "[0.1]"))
    sql, args = pool.calls[0][1], pool.calls[0][2]
    assert "c.workspace_id = $1" in sql and "d.status = 'ready'" in sql and args[0] == WS
    block = CD.format_block(rows)
    assert block.startswith("COMPANY DOCUMENTS") and "[Company: TSMS §4.2 Emergency Drills]" in block
    pool = _Pool(fetch=[{"title": "TSMS", "section": "4.2 Emergency Drills", "text": "monthly"}])
    hit = asyncio.run(CD.lookup(pool, WS, "Company: tsms §4.2 Emergency Drills"))
    assert hit == {"title": "TSMS", "section": "4.2 Emergency Drills", "text": "monthly"}
    assert "c.workspace_id = $1" in pool.calls[0][1] and pool.calls[0][2][0] == WS
    assert asyncio.run(CD.lookup(_Pool(), WS, "not a company citation")) is None


def test_context_block_is_none_without_ready_documents():
    pool = _Pool(fetchval=[None])
    assert asyncio.run(CD.context_block(pool, WS, "key", "drills")) is None
    assert len(pool.calls) == 1          # no embedding call, no search


def _upload(name, content_type, data=b"%PDF-1.4 text"):
    return UploadFile(filename=name, file=io.BytesIO(data), headers={"content-type": content_type})


@pytest.fixture
def roles(monkeypatch):
    state = {"role": "owner"}

    async def require(pool, workspace_id, user_id, required):
        if state["role"] is None:
            raise HTTPException(404, "Workspace not found")
        if state["role"] not in required:
            raise HTTPException(403, "role")
        return state["role"]
    monkeypatch.setattr(R, "_require_workspace_role", require)
    return state


def test_upload_rules(roles, tmp_path, monkeypatch):
    monkeypatch.setattr(R.settings, "upload_dir", str(tmp_path))
    queued = []
    import app.tasks as T
    monkeypatch.setattr(T.process_company_document, "delay", lambda doc_id: queued.append(doc_id))
    row = {"id": uuid.uuid4(), "title": "TSMS", "filename": "TSMS.pdf", "mime_type": "application/pdf",
           "size_bytes": 13, "pages": None, "chunk_count": 0, "status": "pending", "error": None,
           "created_at": datetime.now(timezone.utc)}

    def run(upload, fetchval=("trialing", 0)):
        return asyncio.run(R.upload_document(WS, USER, _Pool(fetchval=fetchval, fetchrow=row), upload))

    assert run(_upload("TSMS.pdf", "application/pdf")).status == "pending" and len(queued) == 1
    saved = list((tmp_path / "workspace_docs" / str(WS)).iterdir())
    assert len(saved) == 1 and saved[0].suffix == ".pdf"
    # .md arrives as octet-stream from browsers
    run(_upload("procedures.md", "application/octet-stream", b"# Drills"))
    with pytest.raises(HTTPException) as e:
        run(_upload("photo.jpg", "image/jpeg"))
    assert e.value.status_code == 422
    with pytest.raises(HTTPException) as e:
        run(_upload("TSMS.pdf", "application/pdf"), fetchval=("archived", 0))
    assert e.value.status_code == 403
    with pytest.raises(HTTPException) as e:
        run(_upload("TSMS.pdf", "application/pdf"), fetchval=("active", CD.MAX_DOCS_PER_WORKSPACE))
    assert e.value.status_code == 422
    roles["role"] = "member"
    with pytest.raises(HTTPException) as e:
        run(_upload("TSMS.pdf", "application/pdf"))
    assert e.value.status_code == 403


def test_members_can_open_a_citation_non_members_cannot(roles, monkeypatch):
    async def fake_lookup(pool, workspace_id, ref):
        return {"title": "TSMS", "section": "4.2 Emergency Drills", "text": "monthly"}
    monkeypatch.setattr(R.company_docs, "lookup", fake_lookup)
    roles["role"] = "member"
    out = asyncio.run(R.open_citation(WS, USER, _Pool(), ref="Company: TSMS §4.2 Emergency Drills"))
    assert out.text == "monthly"
    roles["role"] = None
    with pytest.raises(HTTPException) as e:
        asyncio.run(R.open_citation(WS, USER, _Pool(), ref="Company: TSMS §4.2"))
    assert e.value.status_code == 404
